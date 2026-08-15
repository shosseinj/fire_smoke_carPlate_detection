from __future__ import annotations

from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
import io
import re
import threading
import queue
from typing import Any, Callable
from zoneinfo import ZoneInfo

import jdatetime
import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, Response, StreamingResponse
from openpyxl.styles import Alignment
from pydantic import BaseModel, Field, model_validator

from app.config import settings
from app.core.auth import require_permission
from app.api.import_progress import ImportJobAccepted
from app.core.attendance_summary_service import AttendanceSummaryService, TimePeriod
from app.core.detection_log_store import DetectionLogRecord, DetectionLogStore
from app.core.detection_media import (
    DetectionMediaStorage,
    InvalidMediaKey,
    MEDIA_STATUS_READY,
)
from app.core.jalali_utils import (
    _get_tehran_tz,
    gregorian_to_jalali_str,
    jalali_month_utc_range,
    local_date_range_bounds_utc,
    local_day_utc_range,
    normalize_digits,
    parse_jalali_date,
    parse_jalali_datetime,
)
from app.core.legacy_detection_service import (
    calculate_access,
    legacy_detection_response,
)
from app.core.personnel_store import normalize_national_code


# ── PATCH request schemas ──────────────────────────────────────────────────


class DetectionLogUpdate(BaseModel):
    """Request body for PATCH /{log_id}/person.

    ``person`` assigns the log to a national code (or ``unknown``) while
    ``detection_time`` rewrites the detection moment from a Jalali datetime
    string in local (Asia/Tehran) time, e.g. ``1404-05-17 08:30:00``.
    At least one editable field must be provided.
    """

    person: str | None = None
    detection_time: str | None = None

    @model_validator(mode="after")
    def _require_editable_field(self) -> DetectionLogUpdate:
        if self.person is not None:
            self.person = self.person.strip()
            if not self.person:
                raise ValueError("شخص باید ارسال شود")
        if self.detection_time is not None:
            self.detection_time = self.detection_time.strip()
            if not self.detection_time:
                raise ValueError("زمان تشخیص باید ارسال شود")
        if self.person is None and self.detection_time is None:
            raise ValueError("حداقل یکی از فیلدهای person یا detection_time باید ارسال شود")
        return self


router = APIRouter(prefix="/api/v1/logs", tags=["Detection Logs"])


class DetectionLogResponse(BaseModel):
    id: int
    person: str | None = None
    confidence: float | None = None
    detection_time: str | None = None
    created_at: str | None = None
    updated_at: str | None = None
    face_image_url: str | None = None
    face_thumbnail: str | None = None
    body_image_url: str | None = None
    body_thumbnail: str | None = None
    snapshot_image_url: str | None = None
    video_url: str | None = None
    video_status: str | None = None
    face_video_status: str | None = None
    media_finalized_at: str | None = None
    unknown_faces_path: str | None = None
    log_type: str | None = None
    ref_img_id: int | str | None = None
    full_name: str | None = None
    room_id: int | None = None
    access_granted: bool | None = None
    counts_for_attendance: bool = True
    created_by: int | None = None
    updated_by: int | None = None
    created_by_username: str | None = None
    updated_by_username: str | None = None
    room_name: str | None = None
    camera_name: str | None = None
    section_name: str | None = None
    building_name: str | None = None
    personnel_id: int | None = None
    camera_id: str | int | None = None
    section_id: int | None = None
    building_id: int | None = None

_PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def _excel_text(value: Any) -> str:
    return "" if value is None else str(value).strip().translate(_PERSIAN_DIGITS).replace("\u200c", " ").lower()


def _excel_bool(value: Any, default: bool) -> bool:
    text = _excel_text(value)
    if not text:
        return default
    if text in {"1", "true", "yes", "y", "بله", "بلی"}:
        return True
    if text in {"0", "false", "no", "n", "خیر", "نه"}:
        return False
    raise ValueError("مقدار بولی نامعتبر است")

_detection_media_storage: DetectionMediaStorage | None = None
_FILTER_MEDIA_WORKERS = 16
_MANUAL_DETECTION_CONFIDENCE = 1.0
_filter_media_executor = ThreadPoolExecutor(
    max_workers=_FILTER_MEDIA_WORKERS,
    thread_name_prefix="detection-media-check",
)
_filter_media_gate = threading.BoundedSemaphore(2)


def get_runtime() -> Any:
    from app.main import runtime
    return runtime


def get_detection_log_store() -> DetectionLogStore:
    return get_runtime().detection_log_store


def get_detection_media_storage() -> DetectionMediaStorage:
    global _detection_media_storage
    if _detection_media_storage is None:
        _detection_media_storage = DetectionMediaStorage(settings.saved_media_path)
    return _detection_media_storage


def get_location_store() -> Any:
    return get_runtime().location_store


def get_personnel_store() -> Any:
    return get_runtime().personnel_store


def _get_current_user_id(current_user: Any) -> int | None:
    if isinstance(current_user, dict):
        return current_user.get("id")
    return getattr(current_user, "id", None)


def _choose_reference_image(
    images: list[Any],
) -> Any | None:
    """Choose a stable reference image: primary first, then lowest image id."""
    if not images:
        return None
    return min(
        images,
        key=lambda img: (
            0 if bool(getattr(img, "is_primary", False)) else 1,
            int(getattr(img, "id", 0)),
        ),
    )


def _calculate_old_person_access(
    person: str,
    personnel_id: int | None,
    room_id: int | None,
    location_store: Any,
) -> bool:
    if "unknown" in person.lower():
        return False
    if room_id in (None, 0):
        return True
    room = location_store.get_room(room_id)
    if room is None:
        return False
    room_name = (room.name or "").lower()
    if "general" in room_name or "عمومی" in room_name:
        return True
    if "forbidden" in room_name or "ممنوع" in room_name:
        return False
    return personnel_id is not None and location_store.check_room_access(personnel_id, room_id)


def _resolve_media_path(relative_path: str | None) -> Path | None:
    try:
        return get_detection_media_storage().resolve(relative_path)
    except InvalidMediaKey:
        return None


def _canonical_media_input(value: Any, field_name: str) -> str | None:
    if value in (None, ""):
        return None
    try:
        return get_detection_media_storage().canonical_key(str(value))
    except InvalidMediaKey as exc:
        raise HTTPException(400, f"{field_name}: {exc}") from exc


def _resolve_usernames(
    record: DetectionLogRecord,
) -> tuple[str | None, str | None]:
    ids = sorted(
        {
            int(user_id)
            for user_id in (record.created_by, record.updated_by)
            if user_id is not None
        }
    )
    if not ids:
        return None, None
    placeholders = ", ".join("?" for _ in ids)
    try:
        with get_runtime().database.connection() as connection:
            rows = connection.execute(
                f"SELECT id, username FROM users WHERE id IN ({placeholders})",
                ids,
            ).fetchall()
    except Exception:
        return None, None
    usernames = {int(row["id"]): row.get("username") for row in rows}
    return (
        usernames.get(int(record.created_by)) if record.created_by is not None else None,
        usernames.get(int(record.updated_by)) if record.updated_by is not None else None,
    )


def _resolve_names(
    room_id: int | None,
    camera_id: str | None,
) -> tuple[str | None, str | None, str | None, str | None]:
    room_name = None
    camera_name = None
    section_name = None
    building_name = None
    ls = get_location_store()
    if room_id is not None:
        room = ls.get_room(room_id)
        if room is not None:
            room_name = room.name
            from app.core.location_store import SectionRecord
            if room.section_id is not None:
                sec = ls.get_section(room.section_id)
                if sec is not None:
                    section_name = sec.name
                    if sec.building_id is not None:
                        bld = ls.get_building(sec.building_id)
                        if bld is not None:
                            building_name = bld.name
    if camera_id is not None:
        try:
            cam = get_runtime().registry.get(camera_id)
            if cam is not None:
                camera_name = cam.name
        except Exception:
            pass
    return room_name, camera_name, section_name, building_name


def _resolve_personnel_name(personnel_id: int | None) -> str | None:
    if personnel_id is None:
        return None
    p = get_personnel_store().get(personnel_id)
    if p is None:
        return None
    return f"{p.fname} {p.lname}"


def _filter_response_enrichment(
    records: list[DetectionLogRecord],
) -> dict[int, dict[str, str | None]] | None:
    """Bulk-load display-only relations used by the filtered log response."""
    if not records:
        return {}
    placeholders = ", ".join("?" for _ in records)
    try:
        with get_runtime().database.connection() as connection:
            rows = connection.execute(
                "SELECT d.id, p.id AS personnel_record_id, p.fname, p.lname, "
                "r.name AS room_name, "
                "s.name AS section_name, b.name AS building_name, "
                "created_user.username AS created_by_username, "
                "updated_user.username AS updated_by_username "
                "FROM detection_logs d "
                "LEFT JOIN personnel p ON p.id = d.personnel_id "
                "LEFT JOIN rooms r ON r.id = d.room_id "
                "LEFT JOIN sections s ON s.id = r.section_id "
                "LEFT JOIN buildings b ON b.id = s.building_id "
                "LEFT JOIN users created_user ON created_user.id = d.created_by "
                "LEFT JOIN users updated_user ON updated_user.id = d.updated_by "
                f"WHERE d.id IN ({placeholders})",
                [record.id for record in records],
            ).fetchall()
    except Exception:
        return None

    result: dict[int, dict[str, str | None]] = {}
    for row in rows:
        full_name = None
        if row.get("personnel_record_id") is not None:
            full_name = f"{row.get('fname')} {row.get('lname')}"
        result[int(row["id"])] = {
            "created_by_username": row.get("created_by_username"),
            "updated_by_username": row.get("updated_by_username"),
            "full_name": full_name,
            "room_name": row.get("room_name"),
            "section_name": row.get("section_name"),
            "building_name": row.get("building_name"),
        }
    return result


def _filter_media_availability(
    records: list[DetectionLogRecord],
) -> dict[str, bool]:
    """Check each referenced media key once using bounded parallel filesystem I/O."""
    candidates: set[str] = set()
    for record in records:
        for value in (record.face_image, record.body_image, record.snapshot_image):
            if value:
                candidates.add(str(value))
        if record.video_status == MEDIA_STATUS_READY and record.video:
            candidates.add(str(record.video))
        if (
            record.face_video_status == MEDIA_STATUS_READY
            and record.face_video_or_unknown_faces
        ):
            candidates.add(str(record.face_video_or_unknown_faces))
    if not candidates:
        return {}

    media = get_detection_media_storage()
    ordered = tuple(candidates)
    # At most two filter requests may schedule checks concurrently. The API's
    # 1000-row page cap also bounds each request to at most 5000 media keys.
    with _filter_media_gate:
        results = _filter_media_executor.map(media.exists, ordered)
        return dict(zip(ordered, results))


def _build_response(
    record: DetectionLogRecord,
    include_detail: bool = False,
    *,
    include_face_thumbnail: bool = True,
    enrichment: dict[str, str | None] | None = None,
    media_availability: dict[str, bool] | None = None,
) -> dict:
    if enrichment is None:
        c_user, u_user = _resolve_usernames(record)
        full_name = _resolve_personnel_name(record.personnel_id)
        room_name, camera_name, section_name, building_name = _resolve_names(
            record.room_id, record.camera_id
        )
    else:
        c_user = enrichment.get("created_by_username")
        u_user = enrichment.get("updated_by_username")
        full_name = enrichment.get("full_name")
        room_name = enrichment.get("room_name")
        section_name = enrichment.get("section_name")
        building_name = enrichment.get("building_name")
        camera_name = None
        if record.camera_id is not None:
            try:
                camera = get_runtime().registry.get(record.camera_id)
                if camera is not None:
                    camera_name = camera.name
            except Exception:
                pass
    response = legacy_detection_response(
        record,
        full_name=full_name,
        room_name=room_name,
        camera_name=camera_name,
        section_name=section_name,
        building_name=building_name,
        created_by_username=c_user,
        updated_by_username=u_user,
        include_detail=include_detail,
    )
    response["detection_time"] = response.get("detection_time_jalali")
    response["created_at"] = response.get("created_at_jalali")
    response["updated_at"] = response.get("updated_at_jalali")

    media = get_detection_media_storage()
    def _media_exists(value: str | None) -> bool:
        if not value:
            return False
        if media_availability is None:
            return media.exists(value)
        return bool(media_availability.get(str(value), False))

    face_ready = _media_exists(record.face_image)
    body_ready = _media_exists(record.body_image)
    snapshot_ready = _media_exists(record.snapshot_image)
    video_ready = record.video_status == MEDIA_STATUS_READY and _media_exists(record.video)
    face_video_ready = (
        record.face_video_status == MEDIA_STATUS_READY
        and _media_exists(record.face_video_or_unknown_faces)
    )
    response.update(
        {
            "face_image_url": f"/api/v1/logs/{record.id}/face" if face_ready else None,
            "face_thumbnail": (
                media.thumbnail_data_uri(record.face_thumbnail, record.face_image)
                if include_face_thumbnail
                else None
            ),
            "body_image_url": f"/api/v1/logs/{record.id}/body" if body_ready else None,
            "body_thumbnail": None,
            "snapshot_image_url": (
                f"/api/v1/logs/{record.id}/snapshot" if snapshot_ready else None
            ),
            "video_url": f"/api/v1/logs/{record.id}/video" if video_ready else None,
            "face_video_url": (
                f"/api/v1/logs/{record.id}/face-video" if face_video_ready else None
            ),
            "unknown_faces_path": (
                f"/api/v1/logs/{record.id}/face-video" if face_video_ready else None
            ),
            "video_status": record.video_status,
            "face_video_status": record.face_video_status,
            "media_finalized_at": (
                record.media_finalized_at.isoformat()
                if hasattr(record.media_finalized_at, "isoformat")
                else record.media_finalized_at
            ),
        }
    )
    for field in (
        "detection_time_utc", "detection_time_local", "detection_time_jalali",
        "created_at_utc", "created_at_local", "created_at_jalali",
        "updated_at_utc", "updated_at_local", "updated_at_jalali",
        "face_image", "body_image", "snapshot_image", "reference_image",
        "snapshot_thumbnail", "import_source_parts",
    ):
        response.pop(field, None)
    return response


def _delete_media_files(record: DetectionLogRecord) -> None:
    store = get_detection_log_store()
    storage = get_detection_media_storage()
    candidates = (
        record.face_image,
        record.face_thumbnail,
        record.body_image,
        record.snapshot_image,
        record.video,
        record.face_video_or_unknown_faces,
    )
    unreferenced: list[str] = []
    for raw in candidates:
        try:
            key = storage.canonical_key(raw)
        except InvalidMediaKey:
            continue
        if key and not store.is_media_key_referenced(key):
            unreferenced.append(key)
    storage.delete_many(unreferenced)


def _parse_jalali_time(s: str) -> tuple[int, int, int]:
    """Parse a Jalali clock string (HH:MM[:SS]) with Persian/Arabic digits."""
    s = normalize_digits(s.strip())
    match = re.fullmatch(r"(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?", s)
    if not match:
        raise ValueError(f"فرمت ساعت شمسی نامعتبر: {s!r}")
    hour, minute = int(match.group(1)), int(match.group(2))
    second = int(match.group(3)) if match.group(3) else 0
    if hour > 23 or minute > 59 or second > 59:
        raise ValueError(f"ساعت شمسی خارج از محدوده: {s!r}")
    return hour, minute, second


def _combine_date_time_to_utc_iso(
    date_jalali: str, time_jalali: str
) -> str:
    """Combine a Jalali date and clock string into a UTC ISO timestamp."""
    from_g = parse_jalali_date(date_jalali)
    from_h, from_m, from_s = _parse_jalali_time(time_jalali)
    tehran = _get_tehran_tz()
    local_dt = datetime(
        from_g.year, from_g.month, from_g.day, from_h, from_m, from_s, tzinfo=tehran
    )
    return local_dt.astimezone(timezone.utc).isoformat()


def _period_to_utc_range(
    period: str,
    from_date_jalali: str | None,
    to_date_jalali: str | None,
    from_datetime_jalali: str | None = None,
    to_datetime_jalali: str | None = None,
    from_time_jalali: str | None = None,
    to_time_jalali: str | None = None,
) -> tuple[str | None, str | None]:
    # Explicit Jalali datetime boundaries take precedence over time-only,
    # date-only boundaries and every preset, including a client-supplied or
    # default period value.
    if isinstance(from_datetime_jalali, str) or isinstance(to_datetime_jalali, str):
        if not from_datetime_jalali or not to_datetime_jalali:
            raise HTTPException(
                400,
                "برای بازه ساعت، from_datetime_jalali و to_datetime_jalali الزامی هستند",
            )
        try:
            from_dt = parse_jalali_datetime(from_datetime_jalali)
            to_dt = parse_jalali_datetime(to_datetime_jalali)
        except ValueError:
            # Fall back to time-only values combined with explicit dates.
            if not from_date_jalali or not to_date_jalali:
                raise HTTPException(
                    400,
                    "برای فیلتر ساعت، from_date_jalali و to_date_jalali الزامی هستند",
                )
            try:
                utc_start = _combine_date_time_to_utc_iso(
                    from_date_jalali, from_datetime_jalali
                )
                utc_end = _combine_date_time_to_utc_iso(
                    to_date_jalali, to_datetime_jalali
                )
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
            if utc_start > utc_end:
                raise HTTPException(400, "زمان شروع نباید بعد از زمان پایان باشد")
            return utc_start, utc_end
        if from_dt > to_dt:
            raise HTTPException(400, "زمان شروع نباید بعد از زمان پایان باشد")
        return (
            from_dt.astimezone(timezone.utc).isoformat(),
            to_dt.astimezone(timezone.utc).isoformat(),
        )
    # Time-only Jalali boundaries combine with the explicit date boundaries
    # to form hour-level windows (from_date+from_time .. to_date+to_time).
    if isinstance(from_time_jalali, str) or isinstance(to_time_jalali, str):
        if not from_date_jalali or not to_date_jalali:
            raise HTTPException(
                400,
                "برای فیلتر ساعت، from_date_jalali و to_date_jalali الزامی هستند",
            )
        if not from_time_jalali or not to_time_jalali:
            raise HTTPException(
                400,
                "برای فیلتر ساعت، from_time_jalali و to_time_jalali الزامی هستند",
            )
        try:
            utc_start = _combine_date_time_to_utc_iso(
                from_date_jalali, from_time_jalali
            )
            utc_end = _combine_date_time_to_utc_iso(
                to_date_jalali, to_time_jalali
            )
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        if utc_start > utc_end:
            raise HTTPException(400, "زمان شروع نباید بعد از زمان پایان باشد")
        return utc_start, utc_end
    # Explicit Jalali boundaries take precedence over every preset, including
    # a client-supplied or default period value.
    if from_date_jalali or to_date_jalali:
        period = "custom"
    if period == "today":
        now_tehran = datetime.now(_get_tehran_tz())
        today_local = now_tehran.date()
        utc_start, utc_end = local_day_utc_range(today_local)
        return utc_start.isoformat(), utc_end.isoformat()
    if period == "last_week":
        now_tehran = datetime.now(_get_tehran_tz())
        today_local = now_tehran.date()
        start_local = today_local - timedelta(days=7)
        utc_start, _ = local_day_utc_range(start_local)
        _, utc_end = local_day_utc_range(today_local)
        return utc_start.isoformat(), utc_end.isoformat()
    if period == "last_month":
        now_tehran = datetime.now(_get_tehran_tz())
        today_local = now_tehran.date()
        start_local = today_local - timedelta(days=30)
        utc_start, _ = local_day_utc_range(start_local)
        _, utc_end = local_day_utc_range(today_local)
        return utc_start.isoformat(), utc_end.isoformat()
    if period == "all":
        return None, None
    if period == "custom":
        if not from_date_jalali or not to_date_jalali:
            raise HTTPException(400, "برای بازه سفارشی، from_date_jalali و to_date_jalali الزامی هستند")
        try:
            from_g = parse_jalali_date(from_date_jalali)
            to_g = parse_jalali_date(to_date_jalali)
            utc_start, utc_end = local_date_range_bounds_utc(from_g, to_g)
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        return utc_start.isoformat(), utc_end.isoformat()
    raise HTTPException(400, f"بازه زمانی ناشناخته است: {period!r}")


# ── Static routes (before /{log_id}) ──────────────────────────────────


@router.post("/log", response_model=DetectionLogResponse)
def create_detection_log(
    body: dict[str, Any],
    current_user: dict = Depends(require_permission("application.read")),
) -> dict:
    store = get_detection_log_store()
    ls = get_location_store()

    personnel_id = body.get("personnel_id")
    room_id = body.get("room_id")
    person = str(body.get("person", "Unknown"))
    confidence = _MANUAL_DETECTION_CONFIDENCE
    ref_img_id = body.get("ref_img_id")
    face_image = _canonical_media_input(
        body.get("face_image_path") or body.get("face_image"), "face_image"
    )
    face_thumbnail = _canonical_media_input(
        body.get("face_thumbnail_path") or body.get("face_thumbnail"),
        "face_thumbnail",
    )
    body_image = _canonical_media_input(
        body.get("body_image_path") or body.get("body_image"), "body_image"
    )
    snapshot_image = _canonical_media_input(
        body.get("snapshot_image_path") or body.get("snapshot_image"),
        "snapshot_image",
    )
    video = _canonical_media_input(
        body.get("video_path") or body.get("video"), "video"
    )
    unknown_faces = _canonical_media_input(
        body.get("unknown_faces_path") or body.get("face_video_or_unknown_faces"),
        "face_video_or_unknown_faces",
    )
    log_type = str(body.get("log_type", "real_time"))
    detection_time = body.get("detection_time", "")
    access_granted = bool(body.get("access_granted", False))
    counts_for_attendance = bool(body.get("counts_for_attendance", True))
    camera_id = body.get("camera_id")

    if not detection_time:
        detection_time = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

    if personnel_id is not None and room_id is not None:
        existing = store.find_dedup(
            personnel_id=int(personnel_id),
            room_id=int(room_id),
            detection_time_utc_str=detection_time,
            time_window_seconds=60,
        )
        if existing is not None:
            if confidence >= existing.confidence:
                store.delete(existing.id)
                _delete_media_files(existing)
            else:
                return _build_response(existing, include_detail=True)

    access = calculate_access(personnel_id, room_id, ls)

    user_id = _get_current_user_id(current_user)
    try:
        record = store.create(
            source_system="api",
            personnel_id=personnel_id,
            person=person,
            confidence=confidence,
            detection_time=detection_time,
            ref_img_id=ref_img_id,
            room_id=room_id,
            camera_id=camera_id,
            access_granted=access,
            counts_for_attendance=counts_for_attendance,
            log_type=log_type,
            face_image=face_image,
            face_thumbnail=face_thumbnail,
            body_image=body_image,
            snapshot_image=snapshot_image,
            video=video,
            face_video_or_unknown_faces=unknown_faces,
            video_status=get_detection_media_storage().finalized_video_status(video),
            face_video_status=get_detection_media_storage().finalized_video_status(
                unknown_faces
            ),
            media_finalized_at=(
                datetime.now(timezone.utc).isoformat()
                if video or unknown_faces
                else None
            ),
            created_by=user_id,
        )
    except ValueError as exc:
        raise HTTPException(400, str(exc))

    media = get_detection_media_storage()
    if media.exists(record.face_image) and not media.exists(record.face_thumbnail):
        generated_thumbnail: str | None = None
        try:
            generated_thumbnail = media.create_face_thumbnail(
                record.face_image,
                filename=f"detection_{record.id}_face_thumbnail.jpg",
            )
            if generated_thumbnail:
                updated_record = store.update(
                    record.id, face_thumbnail=generated_thumbnail
                )
                if updated_record is not None:
                    record = updated_record
                else:
                    media.delete(generated_thumbnail)
        except (OSError, RuntimeError):
            if generated_thumbnail:
                media.delete(generated_thumbnail)

    return _build_response(record, include_detail=True)


@router.get("/filter", response_model=list[DetectionLogResponse])
def filter_logs(
    period: str = Query("all"),
    from_date_jalali: str | None = Query(None),
    to_date_jalali: str | None = Query(None),
    from_datetime_jalali: str | None = Query(None),
    to_datetime_jalali: str | None = Query(None),
    from_time_jalali: str | None = Query(None),
    to_time_jalali: str | None = Query(None),
    personnel_id: int | None = Query(None),
    national_code: str | None = Query(None),
    room_id: int | None = Query(None),
    camera_id: str | None = Query(None),
    section_id: int | None = Query(None),
    building_id: int | None = Query(None),
    access_granted: bool | None = Query(None),
    counts_for_attendance: bool | None = Query(None),
    log_type: str | None = Query(None),
    min_confidence: float | None = Query(None),
    max_confidence: float | None = Query(None),
    include_thumbnails: bool = Query(False),
    skip: int = Query(0, ge=0),
    limit: int = Query(200, ge=1, le=1000),
    _: dict = Depends(require_permission("application.read")),
) -> list[dict]:
    from_date_utc, to_date_utc = _period_to_utc_range(
        period,
        from_date_jalali,
        to_date_jalali,
        from_datetime_jalali,
        to_datetime_jalali,
        from_time_jalali,
        to_time_jalali,
    )

    store = get_detection_log_store()
    records, _ = store.list_filter(
        offset=skip,
        limit=limit,
        personnel_id=personnel_id,
        national_code=national_code,
        room_id=room_id,
        camera_id=camera_id,
        section_id=section_id,
        building_id=building_id,
        access_granted=access_granted,
        counts_for_attendance=counts_for_attendance,
        log_type=log_type,
        min_confidence=min_confidence,
        max_confidence=max_confidence,
        from_date_utc=from_date_utc,
        to_date_utc=to_date_utc,
        include_thumbnails=include_thumbnails,
        include_total=False,
    )
    enrichment = _filter_response_enrichment(records)
    media_availability = _filter_media_availability(records)
    return [
        _build_response(
            r,
            include_detail=True,
            include_face_thumbnail=include_thumbnails,
            enrichment=(enrichment.get(r.id, {}) if enrichment is not None else None),
            media_availability=media_availability,
        )
        for r in records
    ]


@router.get("/daily-summary")
def daily_summary(
    period: TimePeriod | None = Query(TimePeriod.TODAY),
    from_date_jalali: str | None = Query(None),
    to_date_jalali: str | None = Query(None),
    include_non_workdays: bool = Query(False),
    include_absent: bool = Query(True),
    _: dict = Depends(require_permission("application.read")),
) -> list[dict]:
    service = AttendanceSummaryService(get_runtime().database)
    return service.daily_summary(
        period=period,
        from_date_jalali=from_date_jalali,
        to_date_jalali=to_date_jalali,
        include_non_workdays=include_non_workdays,
        include_absent=include_absent,
    )


@router.get("/monthly-summary")
def monthly_summary(
    jalali_year: int = Query(..., description="سال شمسی، مثال: ۱۴۰۴"),
    jalali_month: int = Query(..., description="ماه شمسی، ۱ تا ۱۲"),
    personnel_id: str | None = Query(
        None,
        description="شناسه پرسنل در پایگاه داده یا کد ملی",
    ),
    section_id: int | None = Query(None),
    shift_id: int | None = Query(None),
    include_daily_rows: bool = Query(
        False,
        description="بازگرداندن جزئیات روزانه برای هر کارمند",
    ),
    move_days: int = Query(
        10,
        ge=0,
        le=29,
        description=(
            "جابه‌جایی مرز ماه شمسی؛ مقدار ۱۰ بازه را از روز ۲۱ ماه قبل "
            "تا روز ۲۰ ماه جاری محاسبه می‌کند. مقدار صفر بازه عادی ماه را برمی‌گرداند"
        ),
    ),
    _: dict = Depends(require_permission("application.read")),
) -> list[dict]:
    service = AttendanceSummaryService(get_runtime().database)
    return service.monthly_summary(
        jalali_year=jalali_year,
        jalali_month=jalali_month,
        personnel_id=personnel_id,
        section_id=section_id,
        shift_id=shift_id,
        include_daily_rows=include_daily_rows,
        move_days=move_days,
    )


@router.get("/monthly-performance")
def monthly_performance(
    jalali_year: int = Query(..., ge=1400, le=1500),
    jalali_month: int = Query(..., ge=1, le=12),
    section_id: int | None = Query(None),
    _: dict = Depends(require_permission("application.read")),
) -> list[dict]:
    utc_start, utc_end = jalali_month_utc_range(jalali_year, jalali_month)
    store = get_detection_log_store()
    records, _ = store.list_filter(
        section_id=section_id,
        from_date_utc=utc_start.isoformat(),
        to_date_utc=utc_end.isoformat(),
    )
    grouped: dict[int, list[DetectionLogRecord]] = defaultdict(list)
    for r in records:
        if r.personnel_id is not None:
            grouped[r.personnel_id].append(r)

    result: list[dict] = []
    for pid, logs in grouped.items():
        full_name = _resolve_personnel_name(pid)
        avg_conf = sum(l.confidence for l in logs) / len(logs) if logs else 0.0
        attendance_count = sum(1 for l in logs if l.counts_for_attendance)
        score = round(avg_conf * 100, 2)
        result.append({
            "personnel_id": pid,
            "full_name": full_name,
            "total_detections": len(logs),
            "attendance_count": attendance_count,
            "average_confidence": round(avg_conf, 4),
            "score": score,
            "quality": "عالی" if score >= 90 else "خوب" if score >= 70 else "متوسط" if score >= 50 else "ضعیف",
        })
    return result


@router.get("/yearly-leave-summary")
def yearly_leave_summary(
    jalali_year: int = Query(..., description="سال شمسی، مثال: ۱۴۰۵"),
    _: dict = Depends(require_permission("application.read")),
) -> list[dict]:
    service = AttendanceSummaryService(get_runtime().database)
    return service.yearly_leave_summary(jalali_year=jalali_year)


@router.get("/import-excel/template")
def import_excel_template(
    _: dict = Depends(require_permission("application.read")),
):
    import io
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "ورود لاگ‌ها"
    ws.sheet_view.rightToLeft = True
    headers = [
        "کد ملی", "سال", "ماه", "روز", "ساعت", "دقیقه",
        "شناسه اتاق", "دسترسی مجاز", "لحاظ در حضور و غیاب",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
        cell.alignment = Alignment(horizontal="right")
    col_widths = [15, 8, 8, 8, 8, 8, 12, 12, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.append(["0012345678", 1403, 6, 15, 8, 30, "", "1", "1"])

    ws_guide = wb.create_sheet("راهنما")
    guide_lines = [
        "راهنمای واردسازی لاگ‌های تشخیص",
        "",
        "ستون‌ها:",
        "A: کد ملی (۱۰ رقمی، اجباری)",
        "B: سال شمسی (۴ رقمی، اجباری)",
        "C: ماه شمسی (۱-۱۲، اجباری)",
        "D: روز شمسی (۱-۳۱، اجباری)",
        "E: ساعت (۰-۲۳، پیش‌فرض ۰)",
        "F: دقیقه (۰-۵۹، پیش‌فرض ۰)",
        "G: شناسه اتاق (عددی، اختیاری)",
        "H: دسترسی مجاز (0 یا 1، پیش‌فرض 1)",
        "I: لحاظ در حضور و غیاب (0 یا 1، پیش‌فرض 1)",
        "",
        "توجه:",
        "- ردیف اول (سرستون) در واردسازی نادیده گرفته می‌شود",
        "- ردیف‌های خالی رد می‌شوند",
    ]
    ws_guide.append(["راهنمای واردسازی"])
    for i, line in enumerate(guide_lines, 1):
        ws_guide.cell(row=i, column=1, value=line)
    ws_guide.column_dimensions["A"].width = 60
    ws_guide.protection.sheet = True
    ws_guide.protection.set_password("readonly")
    wb.security.lockStructure = True
    wb.security.set_workbook_password("readonly")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=import_template.xlsx"},
    )


def _import_excel_sync(
    file: UploadFile,
    calendar: str = Query("jalali", pattern="^(jalali|gregorian)$"),
    skip_duplicates: bool = Query(True),
    current_user: dict = Depends(require_permission("application.read")),
    progress_callback: Callable[[int, int, int, int], None] | None = None,
) -> dict:
    if file.filename and not (file.filename.endswith(".xlsx") or file.filename.endswith(".xlsm")):
        raise HTTPException(400, "فقط فایل‌های .xlsx یا .xlsm پذیرفته می‌شوند")

    import io
    contents = file.file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents))
    ws = wb.active
    if ws is None:
        raise HTTPException(400, "فایل اکسل برگه فعالی ندارد")

    store = get_detection_log_store()
    ls = get_location_store()
    ps = get_personnel_store()
    user_id = _get_current_user_id(current_user)

    parsed_rows: list[dict[str, Any]] = []
    failed = 0
    skipped = 0
    imported = 0
    failed_row_details: list[dict[str, Any]] = []
    skipped_row_details: list[dict[str, Any]] = []
    successful_rows: list[dict[str, Any]] = []
    excel_rows = [
        (row_idx, row)
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
        if row and any(value is not None for value in row)
    ]
    total_rows = len(excel_rows)

    def emit_progress() -> None:
        if progress_callback is not None:
            progress_callback(total_rows, imported, skipped, failed)

    emit_progress()
    for row_idx, row in excel_rows:
        try:
            national_code = _excel_text(row[0]).replace(" ", "")
            year = int(_excel_text(row[1])) if row[1] is not None else None
            month = int(_excel_text(row[2])) if row[2] is not None else None
            day = int(_excel_text(row[3])) if row[3] is not None else None
            hour = int(_excel_text(row[4])) if row[4] is not None else 0
            minute = int(_excel_text(row[5])) if row[5] is not None else 0
            room_id = int(_excel_text(row[6])) if row[6] not in (None, "") else None
            explicit_access = _excel_bool(row[7], True) if len(row) > 7 else None
            counts_for_attendance = _excel_bool(row[8], True) if len(row) > 8 else True
            if not national_code or year is None or month is None or day is None:
                raise ValueError("missing required import fields")
            if calendar == "gregorian":
                g_dt = datetime(year, month, day, hour, minute)
            else:
                j_dt = jdatetime.datetime(year, month, day, hour, minute)
                g_dt = j_dt.togregorian()
            detection_time = (
                g_dt.replace(tzinfo=_get_tehran_tz())
                .astimezone(timezone.utc)
                .isoformat()
            )
        except (ValueError, TypeError, IndexError) as exc:
            failed += 1
            failed_row_details.append({
                "row": row_idx,
                "national_code": _excel_text(row[0]) if row else "",
                "code": "invalid_row",
                "message": str(exc),
            })
            emit_progress()
            continue
        parsed_rows.append({
            "row_idx": row_idx,
            "national_code": national_code,
            "detection_time": detection_time,
            "room_id": room_id,
            "explicit_access": explicit_access,
            "counts_for_attendance": counts_for_attendance,
        })

    personnel_by_code = ps.get_by_national_codes(
        {str(row["national_code"]) for row in parsed_rows}
    )
    resolved_rows: list[dict[str, Any]] = []
    for row in parsed_rows:
        personnel = personnel_by_code.get(
            normalize_national_code(str(row["national_code"]))
        )
        if personnel is None:
            failed += 1
            failed_row_details.append({
                "row": row["row_idx"],
                "national_code": row["national_code"],
                "code": "personnel_not_found",
                "message": "پرسنل با کد ملی داده‌شده یافت نشد.",
            })
            emit_progress()
            continue
        resolved_rows.append({**row, "personnel": personnel})

    accepted_rows: list[dict[str, Any]] = []
    duplicate_candidate_indexes: set[int] = set()
    candidate_index_by_row: dict[int, int] = {}
    if skip_duplicates:
        candidates: list[tuple[int, int, str]] = []
        for row_position, row in enumerate(resolved_rows):
            personnel = row["personnel"]
            room_id = row["room_id"]
            if personnel.id is None or room_id is None:
                continue
            candidate_index_by_row[row_position] = len(candidates)
            candidates.append((personnel.id, room_id, str(row["detection_time"])))
        duplicate_candidate_indexes = store.find_dedup_indexes(
            candidates, time_window_seconds=60
        )

    accepted_times: dict[tuple[int, int], list[datetime]] = defaultdict(list)
    for row_position, row in enumerate(resolved_rows):
        personnel = row["personnel"]
        room_id = row["room_id"]
        if skip_duplicates and personnel.id is not None and room_id is not None:
            candidate_index = candidate_index_by_row.get(row_position)
            if candidate_index in duplicate_candidate_indexes:
                skipped += 1
                skipped_row_details.append({
                    "row": row["row_idx"],
                    "national_code": row["national_code"],
                    "code": "duplicate_existing",
                })
                emit_progress()
                continue
            current_time = datetime.fromisoformat(str(row["detection_time"]))
            key = (personnel.id, room_id)
            if any(
                previous <= current_time
                and previous >= current_time - timedelta(seconds=60)
                for previous in accepted_times[key]
            ):
                skipped += 1
                skipped_row_details.append({
                    "row": row["row_idx"],
                    "national_code": row["national_code"],
                    "code": "duplicate_in_file",
                })
                emit_progress()
                continue
            accepted_times[key].append(current_time)
        accepted_rows.append(row)

    access_pairs = {
        (row["personnel"].id, row["room_id"])
        for row in accepted_rows
        if row["explicit_access"] is None
        and row["personnel"].id is not None
        and row["room_id"] is not None
    }
    access_by_pair = ls.resolve_access_for_pairs(access_pairs)

    payloads: list[dict[str, Any]] = []
    for row in accepted_rows:
        personnel = row["personnel"]
        room_id = row["room_id"]
        explicit_access = row["explicit_access"]
        if explicit_access is not None:
            access = explicit_access
        elif room_id is None:
            access = True
        else:
            access = access_by_pair.get((personnel.id, room_id), False)
        payloads.append({
            "source_system": "excel_import",
            "personnel_id": personnel.id,
            "person": f"{personnel.fname} {personnel.lname}",
            "confidence": _MANUAL_DETECTION_CONFIDENCE,
            "detection_time": row["detection_time"],
            "room_id": room_id,
            "camera_id": None,
            "access_granted": access,
            "counts_for_attendance": row["counts_for_attendance"],
            "log_type": "excel_import",
            "created_by": user_id,
        })

    for start in range(0, len(payloads), 500):
        batch = payloads[start:start + 500]
        source_batch = accepted_rows[start:start + 500]
        try:
            records = store.create_many(batch, batch_size=len(batch))
            imported += len(records)
            successful_rows.extend({
                "row": source["row_idx"],
                "national_code": source["national_code"],
                "detection_log_id": record.id,
            } for source, record in zip(source_batch, records))
            emit_progress()
        except Exception:
            # Preserve the previous row-level partial-success behavior if one
            # exceptional record prevents a whole optimized batch insert.
            for payload, source in zip(batch, source_batch):
                try:
                    record = store.create(**payload)
                    imported += 1
                    successful_rows.append({
                        "row": source["row_idx"],
                        "national_code": source["national_code"],
                        "detection_log_id": record.id,
                    })
                    emit_progress()
                except Exception as exc:
                    failed += 1
                    failed_row_details.append({
                        "row": source["row_idx"],
                        "national_code": source["national_code"],
                        "code": "database_insert_failed",
                        "message": type(exc).__name__,
                    })
                    emit_progress()

    return {
        "imported_rows": imported,
        "skipped_rows": skipped,
        "failed_rows": failed,
        "successful_row_details": successful_rows,
        "skipped_row_details": skipped_row_details,
        "failed_row_details": failed_row_details,
    }


@router.post(
    "/import-excel",
    status_code=202,
    response_model=ImportJobAccepted,
)
async def import_excel(
    file: UploadFile,
    calendar: str = Query("jalali", pattern="^(jalali|gregorian)$"),
    skip_duplicates: bool = Query(True),
    current_user: dict = Depends(require_permission("application.read")),
) -> dict[str, Any]:
    filename = file.filename or "detection-logs.xlsx"
    if not filename.casefold().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "فقط فایل‌های .xlsx یا .xlsm پذیرفته می‌شوند")
    contents = await file.read()
    if not contents:
        raise HTTPException(422, "فایل اکسل خالی است")
    if len(contents) > 10 * 1024 * 1024:
        raise HTTPException(413, "حجم فایل اکسل نباید بیشتر از ۱۰ مگابایت باشد.")

    def processor(report):
        upload = UploadFile(filename=filename, file=io.BytesIO(contents))
        result = _import_excel_sync(
            upload,
            calendar=calendar,
            skip_duplicates=skip_duplicates,
            current_user=current_user,
            progress_callback=report,
        )
        total = result["imported_rows"] + result["skipped_rows"] + result["failed_rows"]
        return {
            "success": result["failed_rows"] == 0,
            "filename": filename,
            **result,
            "summary": {
                "total_rows": total,
                "imported": result["imported_rows"],
                "skipped": result["skipped_rows"],
                "failed": result["failed_rows"],
            },
        }

    try:
        runtime = get_runtime()
        record = runtime.excel_imports.submit(
            "detection_logs_excel",
            filename,
            _get_current_user_id(current_user),
            processor,
        )
    except queue.Full:
        raise HTTPException(503, "صف ورود فایل‌های اکسل پر است؛ بعداً دوباره تلاش کنید.")
    return {
        "job_id": record.id,
        "progress_id": record.id,
        "status": record.status,
        "status_url": f"/api/v1/import-progress/{record.id}",
        "message": "فایل دریافت شد و پردازش آن در صف قرار گرفت.",
    }


@router.delete("/delete-all-logs")
def delete_all_logs(
    current_user: dict = Depends(require_permission("application.system")),
) -> dict:
    store = get_detection_log_store()
    media_storage = get_detection_media_storage()
    for media_keys in store.iter_media_key_batches(batch_size=1000):
        media_storage.delete_many(media_keys)
    count = store.delete_all()
    return {"deleted_count": count}


# ── Generate fake detections (admin only) ────────────────────────────


@router.post("/generate-fake", status_code=201)
def generate_fake_detections(
    from_date: str | None = Query(None, description="تاریخ شروع jalali"),
    to_date: str | None = Query(None, description="تاریخ پایان jalali (همراه from_date)"),
    pair_logs: bool = Query(False, description="ایجاد لاگ‌های جفتی (ورود+خروج) با فاصله چند دقیقه"),
    remove_existing: bool = Query(
        False,
        description="حذف همه تشخیص‌های موجود در بازه پیش از تولید",
    ),
    admin_user: Any = Depends(require_permission("application.manage")),
) -> dict[str, Any]:
    """Generate fake detection logs for testing/demo purposes.

    Supports optional Jalali date-range filtering with from_date/to_date,
    automatic personnel/room/camera selection, and paired entry/exit log
    generation. Paired logs use 06:00–09:00 for entry and 14:00–18:00 for
    exit. Non-paired logs use 07:00–18:00 in the business timezone.
    """
    import random as _random

    store = get_detection_log_store()
    ps = get_personnel_store()
    ls = get_location_store()

    # ── Resolve date range ──────────────────────────────────────────
    local_tz = ZoneInfo(settings.business_timezone_name)
    range_start: date | None = None
    range_end: date | None = None

    if bool(from_date) != bool(to_date):
        raise HTTPException(400, "from_date و to_date باید با هم ارسال شوند")
    if remove_existing and not (from_date and to_date):
        raise HTTPException(400, "برای حذف تشخیص‌های موجود، from_date و to_date الزامی هستند")
    if from_date and to_date:
        try:
            from_g = parse_jalali_date(from_date)
            to_g = parse_jalali_date(to_date)
            range_start = from_g
            range_end = to_g
        except Exception:
            raise HTTPException(400, "فرمت تاریخ نامعتبر است (jalali: YYYY-MM-DD)")
        if range_start > range_end:
            raise HTTPException(400, "تاریخ شروع نباید بعد از تاریخ پایان باشد")
    # ── Resolve personnel ───────────────────────────────────────────
    selected_personnel: list[Any] = []
    personnel_offset = 0
    personnel_total = 1
    while personnel_offset < personnel_total:
        personnel_page, personnel_total = ps.list(
            offset=personnel_offset,
            limit=1000,
        )
        selected_personnel.extend(personnel_page)
        if not personnel_page:
            break
        personnel_offset += len(personnel_page)
    if not selected_personnel:
        raise HTTPException(404, "هیچ پرسنلی یافت نشد")

    # ── Resolve rooms ────────────────────────────────────────────────
    all_rooms, _ = ls.list_rooms(limit=500)
    selected_room_ids = [r.id for r in all_rooms] if all_rooms else [None]

    # ── Resolve cameras ───────────────────────────────────────────────
    all_cams = get_runtime().registry.list()
    selected_camera_ids = [c.source_uri for c in all_cams] if all_cams else [None]

    deleted_count = 0
    if remove_existing and range_start is not None and range_end is not None:
        utc_start, _ = local_day_utc_range(range_start, local_tz)
        _, utc_end = local_day_utc_range(range_end, local_tz)
        deleted_records = store.delete_in_time_range(
            utc_start.isoformat(), utc_end.isoformat()
        )
        deleted_count = len(deleted_records)
        for record in deleted_records:
            _delete_media_files(record)

    def _random_between(start: datetime, end: datetime) -> datetime:
        """Return a random timezone-aware datetime in an inclusive interval."""
        span_seconds = max(0, int((end - start).total_seconds()))
        return start + timedelta(seconds=_random.randint(0, span_seconds))

    def _window(day: date, start_hour: int, end_hour: int) -> tuple[datetime, datetime]:
        return (
            datetime(day.year, day.month, day.day, start_hour, tzinfo=local_tz),
            datetime(day.year, day.month, day.day, end_hour, tzinfo=local_tz),
        )

    if range_start is None or range_end is None:
        selected_days = [datetime.now(local_tz).date()]
    else:
        selected_days = [
            range_start + timedelta(days=offset)
            for offset in range((range_end - range_start).days + 1)
        ]

    created = 0

    for detection_day in selected_days:
        for person in selected_personnel:
            rid = _random.choice(selected_room_ids)
            cid = _random.choice(selected_camera_ids) if selected_camera_ids else None

            def _create_one(dt: datetime) -> bool:
                if person.id is not None and rid is not None:
                    access = calculate_access(person.id, rid, ls)
                else:
                    access = _random.random() > 0.2
                dt_utc = dt.astimezone(timezone.utc)
                try:
                    store.create(
                        source_system="generate_fake",
                        personnel_id=person.id,
                        person=f"{person.fname} {person.lname}",
                        confidence=1,
                        detection_time=dt_utc.isoformat(),
                        room_id=rid,
                        camera_id=cid,
                        access_granted=access,
                        counts_for_attendance=True,
                        log_type=_random.choice(["camera_rtsp", "tehran_door", "excel_import"]),
                    )
                    return True
                except Exception:
                    return False

            if pair_logs:
                entry_time = _random_between(*_window(detection_day, 6, 9))
                exit_time = _random_between(*_window(detection_day, 14, 18))
                if _create_one(entry_time):
                    created += 1
                if _create_one(exit_time):
                    created += 1
                continue

            daily_count = _random.randint(0, 5)
            day_start, day_end = _window(detection_day, 7, 18)
            detection_times = sorted(
                _random_between(day_start, day_end) for _ in range(daily_count)
            )
            for detection_time in detection_times:
                if _create_one(detection_time):
                    created += 1

    return {
        "message": f"{created} لاگ آزمایشی ایجاد شد",
        "count": created,
        "deleted_count": deleted_count,
    }


# ── Parameterized routes ──────────────────────────────────────────────


@router.get("/{log_id}", response_model=DetectionLogResponse)
def get_log(
    log_id: int,
    _: dict = Depends(require_permission("application.read")),
) -> dict:
    store = get_detection_log_store()
    record = store.get(log_id)
    if record is None:
        raise HTTPException(404, "لاگ یافت نشد")
    return _build_response(record, include_detail=True)


@router.patch("/{log_id}/person", response_model=DetectionLogResponse)
def patch_log_person(
    log_id: int,
    update_data: DetectionLogUpdate,
    current_user: dict = Depends(require_permission("application.manage")),
) -> dict:
    """Assign a detection log to a national code, mark it unknown, or rewrite
    its detection time from a Jalali datetime in local time."""
    store = get_detection_log_store()
    record = store.get(log_id)
    if record is None:
        raise HTTPException(404, "لاگ تشخیص یافت نشد")

    # ── Resolve detection time (Jalali, local time) ───────────────────
    kwargs: dict[str, Any] = {
        "updated_by": _get_current_user_id(current_user),
    }
    if update_data.detection_time:
        try:
            local_dt = parse_jalali_datetime(update_data.detection_time)
            kwargs["detection_time"] = local_dt.astimezone(timezone.utc).isoformat()
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc

    if update_data.person is None:
        updated = store.update(log_id, **kwargs)
        if updated is None:
            raise HTTPException(404, "لاگ تشخیص یافت نشد")
        return _build_response(updated, include_detail=True)

    # ── Resolve current / new personnel ──────────────────────────────
    personnel_store = get_personnel_store()

    person = update_data.person.strip()
    if person.casefold() == "unknown":
        personnel = None
        new_personnel_id = None
        new_ref_img_id = None
        confidence = 0.0
    else:
        personnel = personnel_store.get_by_national_code(person)
        if personnel is None:
            raise HTTPException(status_code=404, detail="کد ملی واردشده در سیستم یافت نشد")
        person = personnel.national_code
        new_personnel_id = personnel.id
        current_personnel_id = record.personnel_id
        if current_personnel_id is None and record.person:
            current = personnel_store.get_by_national_code(record.person)
            current_personnel_id = current.id if current is not None else None
        new_ref_img_id = record.ref_img_id
        if new_personnel_id != current_personnel_id:
            new_ref_img_id = None
            images = personnel_store.list_images(new_personnel_id)
            selected_image = _choose_reference_image(images)
            if selected_image is not None:
                new_ref_img_id = str(selected_image.id)
        confidence = 1.0

    # ── Recalculate access ───────────────────────────────────────────
    location_store = get_location_store()
    access_granted = _calculate_old_person_access(
        person,
        new_personnel_id,
        record.room_id,
        location_store,
    )

    # ── Persist ──────────────────────────────────────────────────────
    kwargs.update({
        "person": person,
        "personnel_id": new_personnel_id,
        "access_granted": access_granted,
        "ref_img_id": new_ref_img_id,
        "confidence": confidence,
    })

    updated = store.update(log_id, **kwargs)
    if updated is None:
        raise HTTPException(404, "لاگ تشخیص یافت نشد")

    return _build_response(updated, include_detail=True)


def _content_disposition(path: Path, download: bool) -> str:
    disposition = "attachment" if download else "inline"
    safe_name = path.name.replace('"', "")
    return f'{disposition}; filename="{safe_name}"'


def _range_not_satisfiable(file_size: int) -> Response:
    return Response(
        status_code=416,
        headers={"Content-Range": f"bytes */{file_size}", "Accept-Ranges": "bytes"},
    )


def _serve_media_file(path: Path, request: Request, download: bool) -> Response:
    media_type = get_detection_media_storage().content_type(path)
    common_headers = {
        "Accept-Ranges": "bytes",
        "Content-Disposition": _content_disposition(path, download),
        "X-Content-Type-Options": "nosniff",
        "Cache-Control": "private, no-store",
    }
    range_header = request.headers.get("range")
    if not range_header:
        return FileResponse(str(path), media_type=media_type, headers=common_headers)

    file_size = path.stat().st_size
    if not range_header.startswith("bytes=") or "," in range_header:
        return _range_not_satisfiable(file_size)
    value = range_header[len("bytes=") :].strip()
    try:
        start_text, end_text = value.split("-", 1)
        if not start_text:
            suffix_length = int(end_text)
            if suffix_length <= 0:
                return _range_not_satisfiable(file_size)
            start = max(0, file_size - suffix_length)
            end = file_size - 1
        else:
            start = int(start_text)
            end = int(end_text) if end_text else file_size - 1
    except (TypeError, ValueError):
        return _range_not_satisfiable(file_size)
    if start < 0 or end < start or start >= file_size:
        return _range_not_satisfiable(file_size)
    end = min(end, file_size - 1)
    content_length = end - start + 1

    def iterator():
        remaining = content_length
        with path.open("rb") as stream:
            stream.seek(start)
            while remaining > 0:
                chunk = stream.read(min(1024 * 1024, remaining))
                if not chunk:
                    break
                remaining -= len(chunk)
                yield chunk

    return StreamingResponse(
        iterator(),
        status_code=206,
        media_type=media_type,
        headers={
            **common_headers,
            "Content-Range": f"bytes {start}-{end}/{file_size}",
            "Content-Length": str(content_length),
        },
    )


def _get_media_path(
    log_id: int,
    kind: str,
) -> tuple[DetectionLogRecord, Path]:
    record = get_detection_log_store().get(log_id)
    if record is None:
        raise HTTPException(404, "لاگ یافت نشد")
    field_by_kind = {
        "face": "face_image",
        "body": "body_image",
        "snapshot": "snapshot_image",
        "video": "video",
        "face-video": "face_video_or_unknown_faces",
    }
    field = field_by_kind[kind]
    if kind == "video" and record.video_status != MEDIA_STATUS_READY:
        status_code = 409 if record.video_status == "writing" else 404
        raise HTTPException(status_code, "ویدیو هنوز آماده نیست")
    if kind == "face-video" and record.face_video_status != MEDIA_STATUS_READY:
        status_code = 409 if record.face_video_status == "writing" else 404
        raise HTTPException(status_code, "ویدیوی چهره هنوز آماده نیست")
    path = _resolve_media_path(getattr(record, field))
    if path is None or not path.is_file():
        raise HTTPException(404, "پرونده یافت نشد")
    return record, path


@router.get("/{log_id}/thumbnail")
@router.get("/{log_id}/media/thumbnail", include_in_schema=False)
def get_thumbnail(
    log_id: int,
    download: bool = Query(False),
    _: dict = Depends(require_permission("application.read")),
) -> Response:
    record = get_detection_log_store().get(log_id)
    if record is None:
        raise HTTPException(404, "لاگ یافت نشد")
    data = get_detection_media_storage().thumbnail_bytes(
        record.face_thumbnail, record.face_image
    )
    if not data:
        raise HTTPException(404, "تصویر بندانگشتی چهره یافت نشد")
    filename = Path(record.face_thumbnail or record.face_image or "face_thumbnail.jpg")
    return Response(
        content=data,
        media_type="image/jpeg",
        headers={
            "Content-Disposition": _content_disposition(filename, download),
            "X-Content-Type-Options": "nosniff",
            "Cache-Control": "private, no-store",
        },
    )


@router.get("/{log_id}/face")
@router.get("/{log_id}/media/face", include_in_schema=False)
def get_face(
    log_id: int,
    request: Request,
    download: bool = Query(False),
    _: dict = Depends(require_permission("application.read")),
) -> Response:
    _, path = _get_media_path(log_id, "face")
    return _serve_media_file(path, request, download)


@router.get("/{log_id}/body")
@router.get("/{log_id}/media/body", include_in_schema=False)
def get_body(
    log_id: int,
    request: Request,
    download: bool = Query(False),
    _: dict = Depends(require_permission("application.read")),
) -> Response:
    _, path = _get_media_path(log_id, "body")
    return _serve_media_file(path, request, download)


@router.get("/{log_id}/snapshot")
@router.get("/{log_id}/media/snapshot", include_in_schema=False)
def get_snapshot(
    log_id: int,
    request: Request,
    download: bool = Query(False),
    _: dict = Depends(require_permission("application.read")),
) -> Response:
    _, path = _get_media_path(log_id, "snapshot")
    return _serve_media_file(path, request, download)


@router.get("/{log_id}/video")
@router.get("/{log_id}/media/video", include_in_schema=False)
def get_video(
    log_id: int,
    request: Request,
    download: bool = Query(False),
    _: dict = Depends(require_permission("application.read")),
) -> Response:
    _, path = _get_media_path(log_id, "video")
    return _serve_media_file(path, request, download)


@router.get("/{log_id}/face-video")
@router.get("/{log_id}/media/face-video", include_in_schema=False)
def get_face_video(
    log_id: int,
    request: Request,
    download: bool = Query(False),
    _: dict = Depends(require_permission("application.read")),
) -> Response:
    _, path = _get_media_path(log_id, "face-video")
    return _serve_media_file(path, request, download)


@router.delete("/{log_id}")
def delete_log(
    log_id: int,
    _: dict = Depends(require_permission("application.system")),
) -> dict:
    store = get_detection_log_store()
    record = store.get(log_id)
    if record is None:
        raise HTTPException(404, "لاگ یافت نشد")
    deleted = store.delete(log_id)
    if not deleted:
        raise HTTPException(404, "لاگ یافت نشد")
    _delete_media_files(record)
    return {
        "message": "لاگ با موفقیت حذف شد",
        "log_id": log_id,
    }
