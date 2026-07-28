"""Official Iranian holiday dataset and Excel import helpers.

Official holidays are supplied as concrete Jalali year/month/day rows. They are
converted to Gregorian ``DATE`` values before being written to PostgreSQL.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, TYPE_CHECKING

if TYPE_CHECKING:
    from app.core.holiday_store import HolidayStore

DEFAULT_OFFICIAL_HOLIDAY_YEAR = 1405
MIN_OFFICIAL_IMPORT_YEAR = 1406
MAX_OFFICIAL_IMPORT_YEAR = 1500
MAX_HOLIDAY_IMPORT_ROWS = 500

# The Excel workbook is intentionally Persian-facing. Internal keys remain
# stable English identifiers so the import service and database are decoupled
# from display labels.
HOLIDAY_TEMPLATE_FIELDS = (
    "month",
    "day",
    "name",
    "description",
    "holiday_type",
)
HOLIDAY_TEMPLATE_HEADERS: dict[str, str] = {
    "month": "ماه",
    "day": "روز",
    "name": "نام تعطیلی",
    "description": "توضیحات",
    "holiday_type": "نوع تعطیلی",
}
HEADER_TO_FIELD = {header: field for field, header in HOLIDAY_TEMPLATE_HEADERS.items()}

HOLIDAY_TYPE_CODE_TO_VALUE: dict[int, str] = {
    1: "national",
    2: "religious",
    3: "company",
    4: "other",
}
HOLIDAY_TYPE_CODE_LABELS: dict[int, str] = {
    1: "ملی",
    2: "مذهبی",
    3: "شرکتی",
    4: "سایر",
}
_DATA_DIRECTORY = Path(__file__).resolve().parents[1] / "data" / "holidays"


@dataclass(frozen=True, slots=True)
class OfficialHolidayRow:
    month: int
    day: int
    name: str
    description: str | None
    holiday_type: str
    every_year: bool
    is_active: bool
    date_value: str

    def to_store_mapping(self) -> dict[str, Any]:
        return {
            "month": self.month,
            "day": self.day,
            "name": self.name,
            "description": self.description,
            "holiday_type": self.holiday_type,
            "every_year": self.every_year,
            "is_active": self.is_active,
            "date_value": self.date_value,
        }


class HolidayImportValidationError(ValueError):
    """Raised when one or more workbook rows fail validation."""

    def __init__(
        self,
        errors: list[dict[str, Any]],
        *,
        total_rows: int,
        valid_rows: int,
    ) -> None:
        super().__init__("Holiday import validation failed")
        self.errors = errors
        self.total_rows = total_rows
        self.valid_rows = valid_rows
        self.invalid_rows = len(errors)

    def response_detail(self, *, year: int, filename: str) -> dict[str, Any]:
        return {
            "message": "فایل اکسل دارای خطاست؛ هیچ تعطیلی‌ای وارد یا حذف نشد.",
            "year": year,
            "filename": filename,
            "database_changed": False,
            "imported_count": 0,
            "summary": {
                "total_data_rows": self.total_rows,
                "valid_rows": self.valid_rows,
                "invalid_rows": self.invalid_rows,
            },
            "choice_codes": excel_choice_code_help(),
            "failed_rows": self.errors,
        }


def excel_choice_code_help() -> dict[str, Any]:
    """Return the numeric choices shown in the Excel guide and API errors."""

    return {
        HOLIDAY_TEMPLATE_HEADERS["holiday_type"]: {
            str(code): label for code, label in HOLIDAY_TYPE_CODE_LABELS.items()
        },
    }


def _validate_dataset_year(year: int) -> int:
    if not DEFAULT_OFFICIAL_HOLIDAY_YEAR <= year <= MAX_OFFICIAL_IMPORT_YEAR:
        raise ValueError(
            "سال جلالی مجموعه تعطیلات باید بین "
            f"{DEFAULT_OFFICIAL_HOLIDAY_YEAR} و {MAX_OFFICIAL_IMPORT_YEAR} باشد."
        )
    return year


def validate_excel_import_year(year: int) -> int:
    """Validate the year accepted by the public Excel-upload endpoint."""

    if not MIN_OFFICIAL_IMPORT_YEAR <= year <= MAX_OFFICIAL_IMPORT_YEAR:
        raise ValueError(
            "سال جلالی برای بارگذاری اکسل باید بین "
            f"{MIN_OFFICIAL_IMPORT_YEAR} و {MAX_OFFICIAL_IMPORT_YEAR} باشد."
        )
    return year


def _normalize_digits(value: str) -> str:
    from app.core.jalali_utils import normalize_digits

    return normalize_digits(value)


def _parse_integer(value: Any, *, field: str) -> int:
    if isinstance(value, bool):
        raise ValueError("باید عدد صحیح باشد.")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        raise ValueError("باید عدد صحیح باشد.")
    text = _normalize_digits(str(value).strip()) if value is not None else ""
    if not text:
        raise ValueError("الزامی است.")
    try:
        numeric = float(text)
    except ValueError as exc:
        raise ValueError("باید عدد صحیح باشد.") from exc
    if not numeric.is_integer():
        raise ValueError("باید عدد صحیح باشد.")
    return int(numeric)


def _parse_internal_boolean(value: Any, *, field: str) -> bool:
    """Parse canonical JSON values used by the bundled 1405 dataset."""

    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in {0, 1}:
        return bool(value)
    text = str(value).strip().casefold() if value is not None else ""
    if text in {"1", "true"}:
        return True
    if text in {"0", "false"}:
        return False
    raise ValueError(f"{field} must be boolean or 0/1")


def _json_safe_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _display_values(raw: Mapping[str, Any]) -> dict[str, Any]:
    return {
        HOLIDAY_TEMPLATE_HEADERS[field]: _json_safe_value(raw.get(field))
        for field in HOLIDAY_TEMPLATE_FIELDS
    }


def _field_error(
    *,
    field: str,
    code: str,
    message: str,
    value: Any,
) -> dict[str, Any]:
    return {
        "field": field,
        "column": HOLIDAY_TEMPLATE_HEADERS.get(field, field),
        "code": code,
        "value": _json_safe_value(value),
        "message": message,
    }


def validate_official_holiday_rows(
    year: int,
    raw_rows: Iterable[Mapping[str, Any]],
    *,
    numeric_excel_choices: bool = False,
) -> list[OfficialHolidayRow]:
    """Validate all rows and return normalized records.

    Validation is all-or-nothing. Every row error is collected before a
    :class:`HolidayImportValidationError` is raised. ``numeric_excel_choices``
    is enabled only for user-uploaded workbooks; the bundled JSON dataset uses
    canonical string/boolean values.
    """

    from app.core.holiday_store import VALID_HOLIDAY_TYPES
    from app.core.jalali_utils import parse_jalali_date

    _validate_dataset_year(year)
    errors: list[dict[str, Any]] = []
    normalized: list[OfficialHolidayRow] = []
    seen: dict[tuple[int, int, str], int] = {}
    rows = list(raw_rows)

    for position, raw in enumerate(rows, start=2):
        row_number_value = raw.get("_row_number", position)
        try:
            row_number = int(row_number_value)
        except (TypeError, ValueError):
            row_number = position
        row_errors: list[dict[str, Any]] = []

        try:
            month = _parse_integer(raw.get("month"), field="month")
            if not 1 <= month <= 12:
                raise ValueError("باید عددی بین 1 و 12 باشد.")
        except ValueError as exc:
            month = 0
            row_errors.append(_field_error(
                field="month",
                code="invalid_month",
                message=str(exc),
                value=raw.get("month"),
            ))

        try:
            day = _parse_integer(raw.get("day"), field="day")
            if not 1 <= day <= 31:
                raise ValueError("باید عددی بین 1 و 31 باشد.")
        except ValueError as exc:
            day = 0
            row_errors.append(_field_error(
                field="day",
                code="invalid_day",
                message=str(exc),
                value=raw.get("day"),
            ))

        name_value = raw.get("name")
        name = str(name_value).strip() if name_value is not None else ""
        if not name:
            row_errors.append(_field_error(
                field="name",
                code="required",
                message="نام تعطیلی الزامی است.",
                value=name_value,
            ))
        elif len(name) < 2:
            row_errors.append(_field_error(
                field="name",
                code="too_short",
                message="نام تعطیلی باید حداقل 2 نویسه باشد.",
                value=name_value,
            ))
        elif len(name) > 255:
            row_errors.append(_field_error(
                field="name",
                code="too_long",
                message="نام تعطیلی نباید بیش از 255 نویسه باشد.",
                value=name_value,
            ))

        description_value = raw.get("description")
        description = None
        if description_value is not None and str(description_value).strip():
            description = str(description_value).strip()
            if len(description) > 4000:
                row_errors.append(_field_error(
                    field="description",
                    code="too_long",
                    message="توضیحات نباید بیش از 4000 نویسه باشد.",
                    value=description_value,
                ))

        holiday_type_value = raw.get("holiday_type")
        holiday_type = ""
        if numeric_excel_choices:
            try:
                holiday_type_code = _parse_integer(
                    holiday_type_value, field="holiday_type"
                )
                holiday_type = HOLIDAY_TYPE_CODE_TO_VALUE[holiday_type_code]
            except KeyError:
                row_errors.append(_field_error(
                    field="holiday_type",
                    code="invalid_choice",
                    message="فقط کدهای 1، 2، 3 یا 4 مجاز هستند.",
                    value=holiday_type_value,
                ))
            except ValueError as exc:
                row_errors.append(_field_error(
                    field="holiday_type",
                    code="required_or_invalid",
                    message=str(exc),
                    value=holiday_type_value,
                ))
        else:
            holiday_type = str(holiday_type_value or "").strip().casefold()
            if holiday_type not in VALID_HOLIDAY_TYPES:
                row_errors.append(_field_error(
                    field="holiday_type",
                    code="invalid_choice",
                    message=(
                        "holiday_type must be one of "
                        f"{sorted(VALID_HOLIDAY_TYPES)}"
                    ),
                    value=holiday_type_value,
                ))

        if numeric_excel_choices:
            # Excel imports always represent concrete official holidays for one
            # Jalali year. They are active when imported and never recurring.
            every_year = False
            is_active = True
        else:
            every_year_value = raw.get("every_year")
            try:
                every_year = _parse_internal_boolean(
                    every_year_value, field="every_year"
                )
                if every_year:
                    raise ValueError(
                        "Official Jalali-year holidays cannot use every_year=true"
                    )
            except ValueError as exc:
                every_year = False
                row_errors.append(_field_error(
                    field="every_year",
                    code="invalid_choice",
                    message=str(exc),
                    value=every_year_value,
                ))

            is_active_value = raw.get("is_active")
            try:
                is_active = _parse_internal_boolean(
                    is_active_value, field="is_active"
                )
            except ValueError as exc:
                is_active = True
                row_errors.append(_field_error(
                    field="is_active",
                    code="invalid_choice",
                    message=str(exc),
                    value=is_active_value,
                ))

        date_value = ""
        if month and day:
            try:
                date_value = parse_jalali_date(
                    f"{year:04d}-{month:02d}-{day:02d}"
                ).isoformat()
            except ValueError as exc:
                row_errors.append(_field_error(
                    field="day",
                    code="invalid_jalali_date",
                    message=f"تاریخ جلالی معتبر نیست: {exc}",
                    value=day,
                ))

        duplicate_key = (month, day, name.casefold())
        if month and day and name:
            previous_row = seen.get(duplicate_key)
            if previous_row is not None:
                row_errors.append(_field_error(
                    field="name",
                    code="duplicate_row",
                    message=(
                        "این تعطیلی با همین ماه، روز و نام تکراری است؛ "
                        f"اولین مورد در ردیف {previous_row} قرار دارد."
                    ),
                    value=name_value,
                ))
            else:
                seen[duplicate_key] = row_number

        if row_errors:
            errors.append({
                "row": row_number,
                "values": _display_values(raw),
                "errors": row_errors,
            })
            continue

        normalized.append(OfficialHolidayRow(
            month=month,
            day=day,
            name=name,
            description=description,
            holiday_type=holiday_type,
            every_year=every_year,
            is_active=is_active,
            date_value=date_value,
        ))

    if not rows:
        errors.append({
            "row": 2,
            "values": {},
            "errors": [{
                "field": "file",
                "column": "فایل",
                "code": "no_data_rows",
                "value": None,
                "message": "فایل اکسل هیچ ردیف تعطیلاتی ندارد.",
            }],
        })

    if errors:
        raise HolidayImportValidationError(
            errors,
            total_rows=len(rows),
            valid_rows=len(normalized),
        )

    normalized.sort(key=lambda item: (item.month, item.day, item.name.casefold()))
    return normalized


def load_default_official_holidays(
    year: int = DEFAULT_OFFICIAL_HOLIDAY_YEAR,
) -> list[OfficialHolidayRow]:
    """Load and validate a bundled official holiday dataset."""

    _validate_dataset_year(year)
    path = _DATA_DIRECTORY / f"iran_{year}.json"
    if not path.is_file():
        raise FileNotFoundError(
            f"No bundled official holiday dataset for Jalali year {year}"
        )
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload_year = int(payload.get("jalali_year", 0))
    if payload_year != year:
        raise ValueError(
            f"Dataset year mismatch: expected {year}, found {payload_year}"
        )
    holidays = payload.get("holidays")
    if not isinstance(holidays, list):
        raise ValueError("Dataset 'holidays' must be a list")
    return validate_official_holiday_rows(year, holidays)


def seed_default_official_holidays(store: "HolidayStore") -> dict[str, int | bool]:
    """Seed official 1405 holidays only when that official year is absent."""

    rows = load_default_official_holidays(DEFAULT_OFFICIAL_HOLIDAY_YEAR)
    return store.seed_official_year_if_missing(
        DEFAULT_OFFICIAL_HOLIDAY_YEAR,
        [row.to_store_mapping() for row in rows],
    )


def build_holiday_excel_template() -> bytes:
    """Create the parameter-free Persian Excel template returned by the API."""

    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "تعطیلات"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    header_border = Border(bottom=thin)

    sheet.append([HOLIDAY_TEMPLATE_HEADERS[field] for field in HOLIDAY_TEMPLATE_FIELDS])
    comments = {
        "month": "ماه جلالی؛ عدد صحیح از 1 تا 12.",
        "day": "روز جلالی معتبر در ماه و سال انتخاب‌شده.",
        "name": "نام مناسبت رسمی؛ حداقل 2 و حداکثر 255 نویسه.",
        "description": "توضیحات اختیاری؛ حداکثر 4000 نویسه.",
        "holiday_type": "کد عددی: 1=ملی، 2=مذهبی، 3=شرکتی، 4=سایر.",
    }
    for field, cell in zip(HOLIDAY_TEMPLATE_FIELDS, sheet[1]):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        cell.comment = Comment(comments[field], "سامانه")
    sheet.row_dimensions[1].height = 26
    widths = {"A": 12, "B": 12, "C": 36, "D": 48, "E": 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    month_validation = DataValidation(
        type="whole", operator="between", formula1="1", formula2="12"
    )
    month_validation.error = "ماه باید یک عدد صحیح بین 1 و 12 باشد."
    month_validation.errorTitle = "ماه نامعتبر"
    month_validation.prompt = "ماه جلالی را با عدد 1 تا 12 وارد کنید."
    month_validation.promptTitle = "ماه جلالی"
    month_validation.showErrorMessage = True
    month_validation.showInputMessage = True
    sheet.add_data_validation(month_validation)
    month_validation.add("A2:A501")

    day_validation = DataValidation(
        type="whole", operator="between", formula1="1", formula2="31"
    )
    day_validation.error = "روز باید یک عدد صحیح بین 1 و 31 باشد."
    day_validation.errorTitle = "روز نامعتبر"
    day_validation.showErrorMessage = True
    sheet.add_data_validation(day_validation)
    day_validation.add("B2:B501")

    name_validation = DataValidation(
        type="textLength", operator="between", formula1="2", formula2="255"
    )
    name_validation.error = "نام تعطیلی باید بین 2 و 255 نویسه باشد."
    name_validation.errorTitle = "نام نامعتبر"
    name_validation.showErrorMessage = True
    sheet.add_data_validation(name_validation)
    name_validation.add("C2:C501")

    description_validation = DataValidation(
        type="textLength", operator="lessThanOrEqual", formula1="4000",
        allow_blank=True,
    )
    description_validation.error = "توضیحات نباید بیش از 4000 نویسه باشد."
    description_validation.errorTitle = "توضیحات طولانی"
    description_validation.showErrorMessage = True
    sheet.add_data_validation(description_validation)
    description_validation.add("D2:D501")

    type_validation = DataValidation(type="list", formula1='"1,2,3,4"')
    type_validation.error = "کد نوع تعطیلی باید یکی از 1، 2، 3 یا 4 باشد."
    type_validation.errorTitle = "نوع تعطیلی نامعتبر"
    type_validation.prompt = "1=ملی، 2=مذهبی، 3=شرکتی، 4=سایر"
    type_validation.promptTitle = "کد نوع تعطیلی"
    type_validation.showErrorMessage = True
    type_validation.showInputMessage = True
    sheet.add_data_validation(type_validation)
    type_validation.add("E2:E501")

    sheet.auto_filter.ref = "A1:E501"

    instructions = workbook.create_sheet("راهنما")
    instructions.sheet_view.rightToLeft = True
    instruction_rows = [
        ["راهنمای ورود تعطیلات رسمی با اکسل"],
        ["سال جلالی در پارامتر year درخواست بارگذاری ارسال می‌شود و نباید ستونی برای سال اضافه شود."],
        ["سال قابل قبول برای بارگذاری", f"از {MIN_OFFICIAL_IMPORT_YEAR} تا {MAX_OFFICIAL_IMPORT_YEAR}"],
        ["ستون‌های الزامی", "ماه، روز، نام تعطیلی، نوع تعطیلی"],
        ["ستون اختیاری", "توضیحات"],
        ["ماه", "عدد صحیح 1 تا 12 در تقویم جلالی"],
        ["روز", "روز معتبر جلالی برای ماه و سال ارسال‌شده"],
        ["نام تعطیلی", "حداقل 2 و حداکثر 255 نویسه؛ چند مناسبت متفاوت می‌توانند یک تاریخ مشترک داشته باشند"],
        ["توضیحات", "اختیاری؛ حداکثر 4000 نویسه"],
        ["کد نوع تعطیلی 1", "ملی"],
        ["کد نوع تعطیلی 2", "مذهبی"],
        ["کد نوع تعطیلی 3", "شرکتی"],
        ["کد نوع تعطیلی 4", "سایر"],
        ["وضعیت رکوردهای واردشده", "تمام تعطیلات بارگذاری‌شده به‌صورت فعال ذخیره می‌شوند."],
        ["تکرار سالانه", "تمام تعطیلات بارگذاری‌شده مخصوص همان سال هستند و با مقدار تکرار سالانه=false ذخیره می‌شوند."],
        ["رفتار اعتبارسنجی", "اگر حتی یک ردیف خطا داشته باشد، کل فایل رد می‌شود و هیچ تغییری در پایگاه داده انجام نمی‌شود."],
        ["رفتار جایگزینی", "فقط تعطیلات رسمی واردشده برای همان سال حذف و جایگزین می‌شوند؛ تعطیلات دستی حفظ می‌شوند."],
        ["نمونه", "ماه=1، روز=1، نام تعطیلی=آغاز نوروز، نوع تعطیلی=1"],
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions["A1"].font = Font(bold=True, size=14, color="1F4E78")
    instructions.column_dimensions["A"].width = 32
    instructions.column_dimensions["B"].width = 100
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    instructions.protection.sheet = True
    instructions.protection.set_password("readonly")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _header_validation_error(
    *,
    header_values: list[str],
    errors: list[dict[str, Any]],
) -> HolidayImportValidationError:
    return HolidayImportValidationError(
        [{
            "row": 1,
            "values": {"سرستون‌ها": header_values},
            "errors": errors,
        }],
        total_rows=0,
        valid_rows=0,
    )


def parse_holiday_excel(data: bytes, *, year: int) -> list[OfficialHolidayRow]:
    """Parse and validate an uploaded Persian Excel workbook without DB writes."""

    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    validate_excel_import_year(year)
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise ValueError(f"فایل اکسل معتبر نیست: {exc}") from exc

    sheet = workbook["تعطیلات"] if "تعطیلات" in workbook.sheetnames else workbook.active
    if sheet is None:
        raise ValueError("فایل اکسل برگه فعال ندارد.")

    header_values = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in sheet[1]
    ]
    header_positions: dict[str, int] = {}
    duplicate_headers: list[str] = []
    unknown_headers: list[str] = []
    for index, header in enumerate(header_values):
        if not header:
            continue
        field = HEADER_TO_FIELD.get(header)
        if field is None:
            unknown_headers.append(header)
            continue
        if field in header_positions:
            duplicate_headers.append(header)
        else:
            header_positions[field] = index

    header_errors: list[dict[str, Any]] = []
    if duplicate_headers:
        header_errors.append({
            "field": "headers",
            "column": "سرستون‌ها",
            "code": "duplicate_headers",
            "value": sorted(set(duplicate_headers)),
            "message": "سرستون تکراری در فایل وجود دارد.",
        })
    if unknown_headers:
        header_errors.append({
            "field": "headers",
            "column": "سرستون‌ها",
            "code": "unknown_headers",
            "value": sorted(set(unknown_headers)),
            "message": "سرستون ناشناخته وجود دارد؛ فقط از قالب دریافت‌شده استفاده کنید.",
        })
    # The template contract requires all five Persian headers. The description
    # cell itself may be blank, but its column must still be present.
    missing_fields = [
        field for field in HOLIDAY_TEMPLATE_FIELDS if field not in header_positions
    ]
    if missing_fields:
        header_errors.append({
            "field": "headers",
            "column": "سرستون‌ها",
            "code": "missing_headers",
            "value": [HOLIDAY_TEMPLATE_HEADERS[field] for field in missing_fields],
            "message": "یک یا چند سرستون الزامی وجود ندارد.",
        })
    if header_errors:
        raise _header_validation_error(
            header_values=header_values,
            errors=header_errors,
        )

    raw_rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        raw = {
            field: values[position] if position < len(values) else None
            for field, position in header_positions.items()
        }
        if all(
            value is None or (isinstance(value, str) and not value.strip())
            for value in raw.values()
        ):
            continue
        raw["_row_number"] = row_number
        raw_rows.append(raw)
        if len(raw_rows) > MAX_HOLIDAY_IMPORT_ROWS:
            raise ValueError(
                f"تعداد ردیف‌ها بیشتر از حد مجاز {MAX_HOLIDAY_IMPORT_ROWS} است."
            )

    return validate_official_holiday_rows(
        year,
        raw_rows,
        numeric_excel_choices=True,
    )
